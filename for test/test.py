from openpyxl import load_workbook

file = "old_1k.xlsx"

wb = load_workbook(file)

schedule = {}
search_text = 'время'

for_schedule = wb.worksheets[:-1]
print(len(for_schedule))
for sheet in for_schedule:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None and search_text in str(cell.value).lower():
                    print( sheet.title, cell.coordinate, cell.value)